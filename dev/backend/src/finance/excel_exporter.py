import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.config.company_identity import (
    COMPANY_LEGAL_NAME,
    COMPANY_ADDRESS,
    COMPANY_TAX_CODE,
    COMPANY_PHONE,
)

# Palette Bách Khoa ERP
COLOR_PRIMARY_DARK = "1E293B"   # Slate 800
COLOR_BRAND_ORANGE = "EB4A23"   # Orange Accent
COLOR_HEADER_BG = "0F172A"      # Header dark background
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_ZEBRA_ROW = "F8FAFC"
COLOR_BORDER = "CBD5E1"
COLOR_MUTED_TEXT = "64748B"

FONT_NAME = "Segoe UI"

FONT_COMPANY = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_PRIMARY_DARK)
FONT_COMPANY_SUB = Font(name=FONT_NAME, size=9, color=COLOR_MUTED_TEXT)
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True, color=COLOR_BRAND_ORANGE)
FONT_SUBTITLE = Font(name=FONT_NAME, size=10, italic=True, color=COLOR_MUTED_TEXT)
FONT_SECTION = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_PRIMARY_DARK)

FONT_TH = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_HEADER_TEXT)
FILL_TH = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")

FONT_TD = Font(name=FONT_NAME, size=10, color=COLOR_PRIMARY_DARK)
FONT_TOTAL = Font(name=FONT_NAME, size=10, bold=True, color=COLOR_PRIMARY_DARK)

FILL_ZEBRA = PatternFill(start_color=COLOR_ZEBRA_ROW, end_color=COLOR_ZEBRA_ROW, fill_type="solid")
FILL_TOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER)
)

BORDER_TOTAL = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="double", color=COLOR_PRIMARY_DARK)
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

NUM_FORMAT_CURRENCY = '#,##0 "₫"'
NUM_FORMAT_INT = '#,##0'

COMPANY_NAME = COMPANY_LEGAL_NAME.upper()
COMPANY_ADDR = f"Địa chỉ: {COMPANY_ADDRESS} | MST: {COMPANY_TAX_CODE} | ĐT: {COMPANY_PHONE}"


def _apply_company_header(ws, title: str, subtitle: str, max_col: int = 7) -> int:
    """Tạo phần tiêu đề công ty và tên báo cáo chuẩn ERP."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell_comp = ws.cell(row=1, column=1, value=COMPANY_NAME)
    cell_comp.font = FONT_COMPANY
    cell_comp.alignment = ALIGN_LEFT

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    cell_sub = ws.cell(row=2, column=1, value=COMPANY_ADDR)
    cell_sub.font = FONT_COMPANY_SUB
    cell_sub.alignment = ALIGN_LEFT

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max_col)
    cell_title = ws.cell(row=4, column=1, value=title.upper())
    cell_title.font = FONT_TITLE
    cell_title.alignment = Alignment(horizontal="center", vertical="center")

    if subtitle:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=max_col)
        cell_subt = ws.cell(row=5, column=1, value=subtitle)
        cell_subt.font = FONT_SUBTITLE
        cell_subt.alignment = Alignment(horizontal="center", vertical="center")
        return 7
    return 6


def _auto_column_width(ws):
    """Tự động co giãn độ rộng các cột dựa trên nội dung."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            # Tránh đo dòng merge cells quá dài
            if cell.row in (1, 2, 4, 5):
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_monthly_dashboard_excel(data: dict, month_str: str) -> io.BytesIO:
    """
    Xuất Báo Cáo Dòng Tiền & Thu Chi Tháng sang file Excel (.xlsx).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bao_Cao_Thang"
    ws.views.sheetView[0].showGridLines = True

    total_income = data.get("total_income", 0)
    total_expenditure = data.get("total_expenditure", 0)
    net_diff = data.get("net_difference", total_income - total_expenditure)
    categories = data.get("categories", [])
    departments = data.get("departments", [])

    cur_row = _apply_company_header(
        ws,
        title=f"BÁO CÁO DÒNG TIỀN VÀ KẾT QUẢ THU CHI - THÁNG {month_str}",
        subtitle=f"Thời gian xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Phân hệ Tài Chính",
        max_col=6
    )

    # 1. BẢNG TỔNG QUAN THU CHI
    ws.cell(row=cur_row, column=1, value="I. TỔNG QUAN DÒNG TIỀN").font = FONT_SECTION
    cur_row += 1

    summary_headers = ["Chỉ tiêu", "Số tiền (VNĐ)", "Tỷ trọng / Ghi chú"]
    for col_idx, h in enumerate(summary_headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = FILL_TH
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER
    cur_row += 1

    summary_rows = [
        ("Tổng Thu Nhập Thực Tế", total_income, "100% doanh thu phát sinh"),
        ("Tổng Chi Phí Hoạt Động", total_expenditure, f"{(total_expenditure / total_income * 100):.1f}% so với thu" if total_income else "—"),
        ("Chênh Lệch Dòng Tiền Thuần (Net)", net_diff, "Thu nhập thặng dư trong kỳ"),
    ]

    for label, amt, note in summary_rows:
        c1 = ws.cell(row=cur_row, column=1, value=label)
        c2 = ws.cell(row=cur_row, column=2, value=amt)
        c3 = ws.cell(row=cur_row, column=3, value=note)

        c1.font = FONT_TD
        c1.border = BORDER_THIN
        c1.alignment = ALIGN_LEFT

        c2.font = FONT_TOTAL if "Chênh Lệch" in label else FONT_TD
        c2.number_format = NUM_FORMAT_CURRENCY
        c2.border = BORDER_THIN
        c2.alignment = ALIGN_RIGHT

        c3.font = FONT_TD
        c3.border = BORDER_THIN
        c3.alignment = ALIGN_LEFT
        cur_row += 1

    cur_row += 2

    # 2. CƠ CẤU THEO DANH MỤC THU / CHI
    ws.cell(row=cur_row, column=1, value="II. CHI TIẾT THEO DANH MỤC THU & CHI").font = FONT_SECTION
    cur_row += 1

    cat_headers = ["STT", "Danh mục", "Tổng Thu (VNĐ)", "Tổng Chi (VNĐ)", "Chênh Lệch (VNĐ)"]
    for col_idx, h in enumerate(cat_headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = FILL_TH
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER
    cur_row += 1

    stt = 1
    for cat in categories:
        c_inc = cat.get("income", 0)
        c_exp = cat.get("expenditure", cat.get("expense", 0))
        c_net = c_inc - c_exp

        r_stt = ws.cell(row=cur_row, column=1, value=stt)
        r_name = ws.cell(row=cur_row, column=2, value=cat.get("name", "Khác"))
        r_inc = ws.cell(row=cur_row, column=3, value=c_inc)
        r_exp = ws.cell(row=cur_row, column=4, value=c_exp)
        r_net = ws.cell(row=cur_row, column=5, value=c_net)

        for cell in (r_stt, r_name, r_inc, r_exp, r_net):
            cell.font = FONT_TD
            cell.border = BORDER_THIN
            if stt % 2 == 0:
                cell.fill = FILL_ZEBRA

        r_stt.alignment = ALIGN_CENTER
        r_name.alignment = ALIGN_LEFT
        r_inc.alignment = ALIGN_RIGHT
        r_exp.alignment = ALIGN_RIGHT
        r_net.alignment = ALIGN_RIGHT

        r_inc.number_format = NUM_FORMAT_CURRENCY
        r_exp.number_format = NUM_FORMAT_CURRENCY
        r_net.number_format = NUM_FORMAT_CURRENCY

        stt += 1
        cur_row += 1

    # Dòng tổng cộng danh mục
    t_stt = ws.cell(row=cur_row, column=1, value="")
    t_name = ws.cell(row=cur_row, column=2, value="TỔNG CỘNG")
    t_inc = ws.cell(row=cur_row, column=3, value=total_income)
    t_exp = ws.cell(row=cur_row, column=4, value=total_expenditure)
    t_net = ws.cell(row=cur_row, column=5, value=net_diff)

    for cell in (t_stt, t_name, t_inc, t_exp, t_net):
        cell.font = FONT_TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL

    t_name.alignment = ALIGN_LEFT
    t_inc.alignment = ALIGN_RIGHT
    t_exp.alignment = ALIGN_RIGHT
    t_net.alignment = ALIGN_RIGHT

    t_inc.number_format = NUM_FORMAT_CURRENCY
    t_exp.number_format = NUM_FORMAT_CURRENCY
    t_net.number_format = NUM_FORMAT_CURRENCY
    cur_row += 3

    # Chữ ký người lập biểu và Giám đốc
    ws.cell(row=cur_row, column=2, value="NGƯỜI LẬP BIỂU").font = FONT_TOTAL
    ws.cell(row=cur_row, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row, column=4, value="GIÁM ĐỐC").font = FONT_TOTAL
    ws.cell(row=cur_row, column=4).alignment = ALIGN_CENTER

    ws.cell(row=cur_row + 1, column=2, value="(Ký, họ tên)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row + 1, column=4, value="(Ký, họ tên, đóng dấu)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=4).alignment = ALIGN_CENTER

    _auto_column_width(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_employee_payroll_excel(ledger_data: dict, month_str: str) -> io.BytesIO:
    """
    Xuất Phiếu Lương Khoán Nhiệm Vụ Cá Nhân sang file Excel (.xlsx).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Phieu_Luong_Ca_Nhan"
    ws.views.sheetView[0].showGridLines = True

    emp = ledger_data.get("employee") or {}
    emp_name = emp.get("full_name", "Nhân viên")
    job_title = emp.get("job_title", "")
    department = emp.get("department", "")
    period_range_val = ledger_data.get("period_range")
    if isinstance(period_range_val, dict):
        period_range_str = period_range_val.get("label") or f"{period_range_val.get('start_date', '')} – {period_range_val.get('end_date', '')}"
    else:
        period_range_str = str(period_range_val or "")
    period_status = ledger_data.get("period_status", "Open")

    summary = ledger_data.get("summary") or {}
    details = ledger_data.get("details") or []
    adjustments = ledger_data.get("adjustments") or []

    cur_row = _apply_company_header(
        ws,
        title=f"PHIẾU LƯƠNG KHOÁN NHIỆM VỤ - KỲ {month_str}",
        subtitle=f"Nhân viên: {emp_name} ({job_title} · {department}) | Kỳ tính lương: {period_range_str}",
        max_col=10
    )

    # 1. TỔNG HỢP THU NHẬP
    ws.cell(row=cur_row, column=1, value="I. TỔNG HỢP THU NHẬP TRONG KỲ").font = FONT_SECTION
    cur_row += 1

    sum_headers = ["Chỉ tiêu", "Số tiền (VNĐ)", "Trạng thái / Ghi chú"]
    for col_idx, h in enumerate(sum_headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = FILL_TH
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER
    cur_row += 1

    approved_amt = summary.get("approved_salary", summary.get("recorded_total", summary.get("approved", 0)))
    pending_amt = summary.get("pending_record_total", summary.get("pending", 0))
    estimated_amt = summary.get("provisional_total", summary.get("estimated_total", summary.get("estimated", 0)))
    allowance_amt = summary.get("allowance", summary.get("total_allowance", 0))
    bonus_amt = summary.get("bonus", summary.get("total_bonus", 0))
    penalty_amt = summary.get("penalty", summary.get("total_deduction", summary.get("total_penalty", 0)))
    net_amt = summary.get("net_salary", summary.get("total_net", summary.get("gross_total", 0)))

    sum_items = [
        ("Lương khoán đã duyệt / Đã ghi nhận", approved_amt, "Đủ điều kiện chi trả"),
        ("Lương khoán chờ ghi nhận (Chờ chốt)", pending_amt, "Nhiệm vụ đã hoàn thành chờ duyệt"),
        ("Lương khoán tạm tính (Chưa xong)", estimated_amt, "Nhiệm vụ đang thực hiện"),
        ("Tổng phụ cấp phát sinh", allowance_amt, "Cọc mốc, hủy hồ sơ, hoàn chi"),
        ("Tổng thưởng phát sinh", bonus_amt, "Thưởng tiến độ, KPI, ưu tiên"),
        ("Tổng phạt / Khấu trừ", penalty_amt, "Khấu trừ nội bộ"),
        ("TỔNG THỰC LĨNH (NET)", net_amt, f"Trạng thái kỳ: {period_status}")
    ]

    for label, amt, note in sum_items:
        c1 = ws.cell(row=cur_row, column=1, value=label)
        c2 = ws.cell(row=cur_row, column=2, value=amt)
        c3 = ws.cell(row=cur_row, column=3, value=note)

        c1.font = FONT_TOTAL if "TỔNG THỰC LĨNH" in label else FONT_TD
        c1.border = BORDER_THIN
        c1.alignment = ALIGN_LEFT

        c2.font = FONT_TOTAL if "TỔNG THỰC LĨNH" in label else FONT_TD
        c2.number_format = NUM_FORMAT_CURRENCY
        c2.border = BORDER_THIN
        c2.alignment = ALIGN_RIGHT

        c3.font = FONT_TD
        c3.border = BORDER_THIN
        c3.alignment = ALIGN_LEFT
        cur_row += 1

    cur_row += 2

    # 2. BẢNG CHI TIẾT CÔNG VIỆC VÀ NHIỆM VỤ
    ws.cell(row=cur_row, column=1, value="II. CHI TIẾT TỪNG NHIỆM VỤ PHÁT SINH LƯƠNG KHOÁN").font = FONT_SECTION
    cur_row += 1

    task_headers = ["STT", "Hồ sơ / Hợp đồng", "Công việc", "Vai trò", "Ngày ghi nhận", "Tiền khoán", "Phụ cấp", "Thưởng / Phạt", "Tổng nhận", "Trạng thái"]
    for col_idx, h in enumerate(task_headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = FILL_TH
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER
    cur_row += 1

    if not details:
        c = ws.cell(row=cur_row, column=1, value="Không có nhiệm vụ phát sinh trong kỳ")
        c.font = FONT_TD
        c.alignment = ALIGN_CENTER
        ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=10)
        for c_idx in range(1, 11):
            ws.cell(row=cur_row, column=c_idx).border = BORDER_THIN
        cur_row += 1
    else:
        stt = 1
        for task in details:
            t_contract = task.get("contract_id") or task.get("contract_code") or task.get("customer_name") or "Hồ sơ nhiệm vụ"
            t_work = task.get("task_name") or task.get("node_name") or "Đo vẽ / Pháp lý"
            t_role = "Chính" if task.get("role") == "main" or task.get("is_primary") else "Phụ"
            t_date = task.get("event_date") or task.get("recorded_at") or task.get("completed_at") or "—"
            t_rate = float(task.get("base_rate", task.get("piece_rate", 0)) or 0)
            t_allow = float((task.get("stake_allowance", 0) or 0) + (task.get("cancellation_allowance", 0) or 0) + (task.get("allowance", 0) or 0))
            t_adj = float((task.get("priority_bonus", task.get("bonus", 0)) or 0) - (task.get("penalty", 0) or 0))
            t_net = float(task.get("net_amount", t_rate + t_allow + t_adj) or 0)
            t_status = task.get("payment_status") or task.get("status", "Đang xử lý")

            cells = [
                ws.cell(row=cur_row, column=1, value=stt),
                ws.cell(row=cur_row, column=2, value=t_contract),
                ws.cell(row=cur_row, column=3, value=t_work),
                ws.cell(row=cur_row, column=4, value=t_role),
                ws.cell(row=cur_row, column=5, value=str(t_date)),
                ws.cell(row=cur_row, column=6, value=t_rate),
                ws.cell(row=cur_row, column=7, value=t_allow),
                ws.cell(row=cur_row, column=8, value=t_adj),
                ws.cell(row=cur_row, column=9, value=t_net),
                ws.cell(row=cur_row, column=10, value=t_status)
            ]

            for cell in cells:
                cell.font = FONT_TD
                cell.border = BORDER_THIN
                if stt % 2 == 0:
                    cell.fill = FILL_ZEBRA

            cells[0].alignment = ALIGN_CENTER
            cells[1].alignment = ALIGN_LEFT
            cells[2].alignment = ALIGN_LEFT
            cells[3].alignment = ALIGN_CENTER
            cells[4].alignment = ALIGN_CENTER
            cells[5].alignment = ALIGN_RIGHT
            cells[6].alignment = ALIGN_RIGHT
            cells[7].alignment = ALIGN_RIGHT
            cells[8].alignment = ALIGN_RIGHT
            cells[9].alignment = ALIGN_CENTER

            cells[5].number_format = NUM_FORMAT_CURRENCY
            cells[6].number_format = NUM_FORMAT_CURRENCY
            cells[7].number_format = NUM_FORMAT_CURRENCY
            cells[8].number_format = NUM_FORMAT_CURRENCY

            stt += 1
            cur_row += 1

    # 3. BẢNG PHỤ CẤP & ĐIỀU CHỈNH LƯƠNG (NẾU CÓ)
    if adjustments:
        cur_row += 2
        ws.cell(row=cur_row, column=1, value="III. CHI TIẾT CÁC KHOẢN ĐIỀU CHỈNH / THƯỞNG / PHỤ CẤP / PHẠT").font = FONT_SECTION
        cur_row += 1

        adj_headers = ["STT", "Ngày áp dụng", "Loại điều chỉnh", "Lý do / Căn cứ", "Số tiền (VNĐ)", "Trạng thái"]
        for col_idx, h in enumerate(adj_headers, 1):
            c = ws.cell(row=cur_row, column=col_idx, value=h)
            c.font = FONT_TH
            c.fill = FILL_TH
            c.border = BORDER_THIN
            c.alignment = ALIGN_CENTER
        cur_row += 1

        adj_labels = {
            "BONUS": "Thưởng (Ưu tiên / KPI)",
            "bonus": "Thưởng",
            "DEDUCTION": "Khấu trừ / Phạt",
            "penalty": "Phạt",
            "ALLOWANCE": "Phụ cấp",
            "allowance": "Phụ cấp",
            "REIMBURSEMENT": "Hoàn chi / Phụ cấp",
            "referral_commission": "Hoa hồng giới thiệu",
            "holiday_bonus": "Thưởng lễ/Tết"
        }

        for a_idx, adj in enumerate(adjustments, 1):
            a_date = adj.get("effective_date") or adj.get("event_date") or "—"
            a_type_raw = str(adj.get("type") or adj.get("adjustment_type") or "Khác")
            a_type_label = adj_labels.get(a_type_raw, a_type_raw)
            a_reason = adj.get("reason") or "Điều chỉnh lương"
            a_amt = float(adj.get("amount") or 0)
            is_neg = a_type_raw.upper() in ("DEDUCTION", "PENALTY")
            if is_neg and a_amt > 0:
                a_amt = -a_amt
            a_status = "Đã duyệt" if adj.get("status") == "approved" else (adj.get("status") or "Đang xử lý")

            cells = [
                ws.cell(row=cur_row, column=1, value=a_idx),
                ws.cell(row=cur_row, column=2, value=str(a_date)),
                ws.cell(row=cur_row, column=3, value=a_type_label),
                ws.cell(row=cur_row, column=4, value=a_reason),
                ws.cell(row=cur_row, column=5, value=a_amt),
                ws.cell(row=cur_row, column=6, value=a_status)
            ]
            for cell in cells:
                cell.font = FONT_TD
                cell.border = BORDER_THIN
                if a_idx % 2 == 0:
                    cell.fill = FILL_ZEBRA
            cells[0].alignment = ALIGN_CENTER
            cells[1].alignment = ALIGN_CENTER
            cells[2].alignment = ALIGN_LEFT
            cells[3].alignment = ALIGN_LEFT
            cells[4].alignment = ALIGN_RIGHT
            cells[4].number_format = NUM_FORMAT_CURRENCY
            cells[5].alignment = ALIGN_CENTER
            cur_row += 1

    cur_row += 3

    # Chữ ký người nhận lương và Kế toán / Giám đốc
    ws.cell(row=cur_row, column=2, value="NGƯỜI NHẬN LƯƠNG").font = FONT_TOTAL
    ws.cell(row=cur_row, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row, column=5, value="KẾ TOÁN").font = FONT_TOTAL
    ws.cell(row=cur_row, column=5).alignment = ALIGN_CENTER
    ws.cell(row=cur_row, column=8, value="GIÁM ĐỐC DUYỆT").font = FONT_TOTAL
    ws.cell(row=cur_row, column=8).alignment = ALIGN_CENTER

    ws.cell(row=cur_row + 1, column=2, value="(Ký, ghi rõ họ tên)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row + 1, column=5, value="(Ký, họ tên)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=5).alignment = ALIGN_CENTER
    ws.cell(row=cur_row + 1, column=8, value="(Ký, họ tên, đóng dấu)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=8).alignment = ALIGN_CENTER

    _auto_column_width(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_department_payroll_summary_excel(department_name: str, period_label: str, summaries: list) -> io.BytesIO:
    """
    Xuất Bảng Lương Khoán Tổng Hợp Phòng Ban sang file Excel (.xlsx).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Bang_Tong_Hop_Luong"
    ws.views.sheetView[0].showGridLines = True

    cur_row = _apply_company_header(
        ws,
        title=f"BẢNG TỔNG HỢP LƯƠNG KHOÁN - PHÒNG BAN: {department_name.upper()}",
        subtitle=f"Kỳ tính lương: {period_label} | Ngày lập bảng: {datetime.now().strftime('%d/%m/%Y')}",
        max_col=8
    )

    headers = [
        "STT", "Họ và Tên", "Chức danh", "Lương Đã Duyệt", "Chờ Ghi Nhận", "Phụ Cấp", "Thưởng / Phạt", "Tổng Thực Lĩnh"
    ]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=cur_row, column=col_idx, value=h)
        c.font = FONT_TH
        c.fill = FILL_TH
        c.border = BORDER_THIN
        c.alignment = ALIGN_CENTER
    cur_row += 1

    stt = 1
    total_approved = 0
    total_pending = 0
    total_allowance = 0
    total_adj = 0
    total_net_all = 0

    for emp in summaries:
        name = emp.get("full_name", "Nhân viên")
        job = emp.get("job_title", "Kỹ thuật viên")
        approved = emp.get("approved", 0)
        pending = emp.get("pending", 0)
        allowance = emp.get("allowance", 0)
        adj = emp.get("bonus", 0) - emp.get("penalty", 0)
        net = emp.get("net_total", approved + allowance + adj)

        total_approved += approved
        total_pending += pending
        total_allowance += allowance
        total_adj += adj
        total_net_all += net

        cells = [
            ws.cell(row=cur_row, column=1, value=stt),
            ws.cell(row=cur_row, column=2, value=name),
            ws.cell(row=cur_row, column=3, value=job),
            ws.cell(row=cur_row, column=4, value=approved),
            ws.cell(row=cur_row, column=5, value=pending),
            ws.cell(row=cur_row, column=6, value=allowance),
            ws.cell(row=cur_row, column=7, value=adj),
            ws.cell(row=cur_row, column=8, value=net),
        ]

        for cell in cells:
            cell.font = FONT_TD
            cell.border = BORDER_THIN
            if stt % 2 == 0:
                cell.fill = FILL_ZEBRA

        cells[0].alignment = ALIGN_CENTER
        cells[1].alignment = ALIGN_LEFT
        cells[2].alignment = ALIGN_LEFT
        for i in range(3, 8):
            cells[i].alignment = ALIGN_RIGHT
            cells[i].number_format = NUM_FORMAT_CURRENCY

        stt += 1
        cur_row += 1

    # Dòng tổng cộng phòng ban
    cells_total = [
        ws.cell(row=cur_row, column=1, value=""),
        ws.cell(row=cur_row, column=2, value="TỔNG CỘNG PHÒNG BAN"),
        ws.cell(row=cur_row, column=3, value=""),
        ws.cell(row=cur_row, column=4, value=total_approved),
        ws.cell(row=cur_row, column=5, value=total_pending),
        ws.cell(row=cur_row, column=6, value=total_allowance),
        ws.cell(row=cur_row, column=7, value=total_adj),
        ws.cell(row=cur_row, column=8, value=total_net_all),
    ]

    for cell in cells_total:
        cell.font = FONT_TOTAL
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL

    cells_total[1].alignment = ALIGN_LEFT
    for i in range(3, 8):
        cells_total[i].alignment = ALIGN_RIGHT
        cells_total[i].number_format = NUM_FORMAT_CURRENCY

    cur_row += 3

    # Chữ ký
    ws.cell(row=cur_row, column=2, value="NGƯỜI LẬP BIỂU").font = FONT_TOTAL
    ws.cell(row=cur_row, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row, column=5, value="KẾ TOÁN TRƯỞNG").font = FONT_TOTAL
    ws.cell(row=cur_row, column=5).alignment = ALIGN_CENTER
    ws.cell(row=cur_row, column=7, value="GIÁM ĐỐC DUYỆT").font = FONT_TOTAL
    ws.cell(row=cur_row, column=7).alignment = ALIGN_CENTER

    ws.cell(row=cur_row + 1, column=2, value="(Ký, họ tên)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=2).alignment = ALIGN_CENTER
    ws.cell(row=cur_row + 1, column=5, value="(Ký, họ tên)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=5).alignment = ALIGN_CENTER
    ws.cell(row=cur_row + 1, column=7, value="(Ký, họ tên, đóng dấu)").font = FONT_SUBTITLE
    ws.cell(row=cur_row + 1, column=7).alignment = ALIGN_CENTER

    _auto_column_width(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
