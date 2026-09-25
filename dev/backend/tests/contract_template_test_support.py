"""Isolated PostgreSQL catalog helpers for contract-template tests.

This module contains helpers only — pytest must not collect it (no ``test_``
prefix, no ``Test*`` names). Import with::

    from contract_template_test_support import isolated_template_catalog, make_docx

``isolated_template_catalog`` builds a private ``tpl_test_<uuid>`` schema on
the already-validated disposable PostgreSQL target, recreates the
``users``/``contract_templates`` baseline (every current column except the four
new management columns), applies the real new migration file, and drops exactly
the schema it created on teardown. It never touches the shared test engine.
"""

from __future__ import annotations

import contextlib
import io
import os
import re
import uuid
from pathlib import Path
from typing import Callable, Iterator, Optional
from urllib.parse import parse_qsl, unquote, urlsplit

import pytest

MIGRATION_FILENAME = "20260925100000_contract_template_management.sql"
SCHEMA_PREFIX = "tpl_test_"
_SCHEMA_RE = re.compile(r"^tpl_test_[0-9a-f]{32}$")

# Local hosts only; "pg-test" is the name of the disposable Postgres service in
# docker-compose.dev.yml, not a public domain. Database names must end with
# _test so a shared/application database is never selected.
_DISPOSABLE_HOSTS = frozenset({"localhost", "127.0.0.1", "::1", "pg-test"})


def make_docx(text="{{customer_name}}"):
    """Build a real minimal DOCX in memory (never an empty ZIP)."""
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def resolve_migration_path() -> Path:
    """Locate the contract-template-management migration file.

    Prefers the container mount (``/app/supabase/migrations``), then falls
    back to the repository checkout relative to this file.
    """
    here = Path(__file__).resolve()
    candidates = [Path("/app/supabase/migrations") / MIGRATION_FILENAME]
    seen = set(candidates)
    for base in [here.parent, *here.parents[:6]]:
        candidate = base / "supabase" / "migrations" / MIGRATION_FILENAME
        if candidate not in seen:
            seen.add(candidate)
            candidates.append(candidate)
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(
        f"contract template migration not found; checked: "
        f"{', '.join(str(c) for c in candidates)}"
    )


def _target_identity(database_url: str):
    parsed = urlsplit(database_url)
    driver = parsed.scheme.lower().split("+", 1)[0]
    host = parsed.hostname.lower() if parsed.hostname else ""
    port = parsed.port or (5432 if driver == "postgresql" else None)
    database = unquote(parsed.path.lstrip("/"))
    query = tuple(sorted(parse_qsl(parsed.query, keep_blank_values=True)))
    return driver, host, port, database, query


def require_disposable_postgres_url() -> str:
    """Return TEST_DATABASE_URL after fail-closed disposable-PG validation.

    Skips (not fails) when no TEST_DATABASE_URL is configured — the caller then
    reports blocked PG tests instead of touching any other database. Raises
    RuntimeError for anything that is explicitly configured but not a local
    disposable ``*_test`` PostgreSQL target.
    """
    dsn = os.environ.get("TEST_DATABASE_URL", "").strip()
    if not dsn:
        pytest.skip(
            "TEST_DATABASE_URL is not configured; contract-template schema "
            "tests require a validated disposable PostgreSQL target"
        )
    try:
        driver, host, _, database, _ = _target_identity(dsn)
    except (TypeError, ValueError) as exc:
        raise RuntimeError("TEST_DATABASE_URL must be a valid database URL.") from exc
    if host.endswith(".supabase.co") or host.endswith(".supabase.com"):
        raise RuntimeError("refusing to use a Supabase database for template tests")
    if (
        driver != "postgresql"
        or host not in _DISPOSABLE_HOSTS
        or not database.endswith("_test")
    ):
        raise RuntimeError(
            "TEST_DATABASE_URL must be a local disposable PostgreSQL *_test "
            "database; refusing to proceed"
        )
    # NOTE: no TEST-vs-DATABASE_URL comparison here. conftest.py already
    # refuses to start when both resolve to the same target, and it then
    # points DATABASE_URL at the validated test target for the whole pytest
    # process — so the two are always equal inside a test run by design.
    return dsn


# Baseline *before* the management migration: every current column of
# contract_templates except the four new ones (upload_state, content_sha256,
# content_size, publish_requested). Executed unqualified under a SET LOCAL
# search_path pointing at the isolated schema.
BASELINE_DDL = """
CREATE TABLE users (
  id VARCHAR(50) PRIMARY KEY,
  username VARCHAR,
  password_hash VARCHAR,
  email VARCHAR,
  is_active BOOLEAN DEFAULT TRUE,
  email_verified BOOLEAN NOT NULL DEFAULT FALSE,
  created_at TIMESTAMPTZ DEFAULT now(),
  updated_at TIMESTAMPTZ DEFAULT now()
);
CREATE TABLE contract_templates (
  id VARCHAR(50) PRIMARY KEY DEFAULT gen_random_uuid()::text,
  code TEXT NOT NULL,
  version INTEGER NOT NULL DEFAULT 1,
  name TEXT NOT NULL,
  description TEXT,
  template_file_name TEXT,
  template_file_link TEXT,
  template_storage_key TEXT,
  storage_provider TEXT NOT NULL DEFAULT 's3-compatible',
  placeholder_schema JSONB NOT NULL DEFAULT '[]'::jsonb,
  render_rules JSONB NOT NULL DEFAULT '{}'::jsonb,
  status TEXT NOT NULL DEFAULT 'draft',
  created_by VARCHAR(50) REFERENCES users(id),
  created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
  updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
  UNIQUE (code, version),
  CONSTRAINT contract_templates_status_check
    CHECK (status IN ('draft', 'published', 'archived')),
  CONSTRAINT contract_templates_storage_provider_check
    CHECK (storage_provider IN ('s3-compatible'))
);
"""


def _drop_schema(dsn: str, schema: str) -> None:
    from sqlalchemy import create_engine
    from sqlalchemy.pool import NullPool

    if not _SCHEMA_RE.match(schema):
        raise RuntimeError(f"refusing to drop unexpected schema {schema!r}")
    killer = create_engine(dsn, poolclass=NullPool)
    try:
        with killer.begin() as conn:
            conn.exec_driver_sql(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
    finally:
        killer.dispose()


class ImmutableStore:
    """Dictionary-backed fake for immutable template storage.

    Enforces immutability like the real provider: an existing key raises
    ``FileExistsError`` and ``put`` always reads the actual bytes. Set
    ``fail_next_put`` to raise a transient error once, then auto-reset.
    ``before_put`` hook may block (threading.Event) to force out-of-order
    completion in concurrency tests. Never used in production.
    """

    def __init__(self):
        self.objects: dict[str, bytes] = {}
        self.fail_next_put: bool = False
        self.before_put = None
        self.put_calls: list[str] = []

    def put(self, file_obj, object_name: str) -> str:
        hook = self.before_put
        if hook is not None:
            hook()
        if self.fail_next_put:
            self.fail_next_put = False
            raise ConnectionError("transient storage failure")
        data = file_obj.read()
        if not isinstance(data, (bytes, bytearray)):
            raise TypeError("fake storage requires bytes")
        data = bytes(data)
        self.put_calls.append(object_name)
        if object_name in self.objects:
            raise FileExistsError(f"Immutable object already exists: {object_name}")
        self.objects[object_name] = data
        return object_name

    def inspect(self, key: str, expected_sha256: str, expected_size: int) -> str:
        stored = self.objects.get(key)
        if stored is None:
            return "missing"
        import hashlib

        if len(stored) == expected_size and hashlib.sha256(stored).hexdigest() == expected_sha256:
            return "matching"
        return "conflict"

    def get(self, object_name: str) -> bytes:
        try:
            return self.objects[object_name]
        except KeyError as exc:
            raise KeyError(f"missing object: {object_name}") from exc


@contextlib.contextmanager
def isolated_template_catalog(
    before_migration: Optional[Callable] = None,
) -> Iterator:
    """Yield a sessionmaker bound to an isolated migrated PG schema.

    ``before_migration(connection)`` optionally seeds legacy rows on the raw
    connection after baseline creation but before the migration runs, so
    backfill/preflight behavior can be tested. Default callback writes nothing.
    """
    from sqlalchemy import event
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import NullPool

    dsn = require_disposable_postgres_url()
    schema = SCHEMA_PREFIX + uuid.uuid4().hex
    migration_sql = resolve_migration_path().read_text(encoding="utf-8")

    engine = create_engine(
        dsn,
        poolclass=NullPool,
        execution_options={
            "schema_translate_map": {None: schema, "public": schema}
        },
    )

    @event.listens_for(engine, "checkout")
    def _pin_search_path(dbapi_conn, conn_record, conn_proxy):
        cursor = dbapi_conn.cursor()
        try:
            cursor.execute(f'SET SESSION search_path = "{schema}"')
        finally:
            cursor.close()

    try:
        with engine.begin() as conn:
            conn.exec_driver_sql(f'CREATE SCHEMA "{schema}"')
            conn.exec_driver_sql(f'SET LOCAL search_path TO "{schema}"')
            conn.exec_driver_sql(BASELINE_DDL)
            if before_migration is not None:
                before_migration(conn)
            # Raw driver cursor on the same DBAPI connection: the migration
            # script is multi-statement and contains PL/pgSQL %-placeholders,
            # which exec_driver_sql would hand to psycopg2 as bound parameters
            # and fail. It still runs inside this transaction.
            raw = conn.connection.driver_connection
            cursor = raw.cursor()
            try:
                cursor.execute(migration_sql)
            finally:
                cursor.close()
    except Exception:
        engine.dispose()
        _drop_schema(dsn, schema)
        raise

    factory = sessionmaker(bind=engine)
    try:
        yield factory
    finally:
        from sqlalchemy.orm import close_all_sessions

        try:
            close_all_sessions()
        except Exception:
            pass
        engine.dispose()
        _drop_schema(dsn, schema)
