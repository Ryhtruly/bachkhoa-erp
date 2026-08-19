from io import BytesIO

import pytest
from botocore.exceptions import ClientError
from docx import Document

from scripts import bootstrap_contract_template


def test_bootstrap_creates_docx_template_only_when_key_is_absent(tmp_path, monkeypatch):
    template = Document()
    template.add_paragraph("{{contract_id}}")
    template_path = tmp_path / "contract.docx"
    template.save(template_path)
    uploaded = {}

    class FakeS3:
        def put_object(self, **kwargs):
            uploaded.update(kwargs)

    monkeypatch.setattr(bootstrap_contract_template, "_get_client", lambda: FakeS3())

    result = bootstrap_contract_template.bootstrap_contract_template(
        template_path,
        "HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx",
    )

    assert result == "contract-templates/HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx"
    assert uploaded["Bucket"] == bootstrap_contract_template.CONTRACT_TEMPLATE_BUCKET
    assert uploaded["Key"] == result
    assert uploaded["Body"].startswith(b"PK")
    assert uploaded["ContentType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert uploaded["IfNoneMatch"] == "*"


def test_bootstrap_maps_atomic_create_conflict_to_no_overwrite_outcome(tmp_path, monkeypatch):
    template = Document()
    template.add_paragraph("{{contract_id}}")
    template_path = tmp_path / "contract.docx"
    template.save(template_path)
    existing_body = b"existing-template"

    class FakeS3:
        def __init__(self):
            self.body = existing_body
            self.requests = []

        def put_object(self, **kwargs):
            self.requests.append(kwargs)
            raise ClientError(
                {"Error": {"Code": "PreconditionFailed"}, "ResponseMetadata": {"HTTPStatusCode": 412}},
                "PutObject",
            )

    client = FakeS3()
    monkeypatch.setattr(bootstrap_contract_template, "_get_client", lambda: client)

    with pytest.raises(FileExistsError, match="contract-templates/HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx"):
        bootstrap_contract_template.bootstrap_contract_template(
            template_path,
            "HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx",
        )

    assert client.body == existing_body
    assert len(client.requests) == 1
    assert client.requests[0]["IfNoneMatch"] == "*"


def test_bootstrap_maps_conditional_request_conflict_to_no_overwrite_outcome(tmp_path, monkeypatch):
    template = Document()
    template.add_paragraph("{{contract_id}}")
    template_path = tmp_path / "contract.docx"
    template.save(template_path)

    class FakeS3:
        def put_object(self, **kwargs):
            raise ClientError(
                {"Error": {"Code": "ConditionalRequestConflict"}, "ResponseMetadata": {"HTTPStatusCode": 409}},
                "PutObject",
            )

    monkeypatch.setattr(bootstrap_contract_template, "_get_client", lambda: FakeS3())

    with pytest.raises(FileExistsError, match="contract-templates/HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx"):
        bootstrap_contract_template.bootstrap_contract_template(
            template_path,
            "HOP_DONG_DICH_VU_KHUNG_BACH_KHOA/v1.docx",
        )
