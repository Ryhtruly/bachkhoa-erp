import io
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from docx import Document

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from src.core import doc_generator
from src.routes import routes_contracts


class ContractDocumentRendererTests(unittest.TestCase):
    def test_document_route_renders_current_persisted_data_not_snapshot(self):
        document = SimpleNamespace(
            output_file_name="HopDong_2004_BK-2026.docx",
            render_data_snapshot={
                "contract_id": "2004/BK-2026",
                "customer_name": "Tên snapshot cũ",
                "customer_email": "snapshot@example.com",
                "_template_version": "mau_hop_dong_v1",
            },
        )
        current_rows = {
            "Contract": SimpleNamespace(
                id="2004/BK-2026",
                customer_id="customer-1",
                lead_id=None,
                contract_template_id=None,
                service_type="Đo hiện trạng",
                total_value=18500000,
                date_signed="2026-08-14",
            ),
            "Customer": SimpleNamespace(
                full_name="Tên hiện tại",
                phone="0900000000",
                address="Địa chỉ hiện tại",
                email="",
            ),
            "ServiceLine": SimpleNamespace(
                service_type="Đo hiện trạng",
                property_address="Địa chỉ hiện tại",
                price=18500000,
            ),
            "Receivable": SimpleNamespace(due_date=None),
            "LeadPipeline": None,
            "ContractGeneratedDocument": document,
        }

        class CurrentDataQuery:
            def __init__(self, row):
                self.row = row

            def filter(self, *_criteria):
                return self

            def order_by(self, *_ordering):
                return self

            def first(self):
                return self.row

        db = SimpleNamespace(query=lambda model: CurrentDataQuery(current_rows.get(model.__name__)))
        rendered = {}
        original_renderer = routes_contracts.doc_generator.render_contract_document
        routes_contracts.doc_generator.render_contract_document = lambda data, version, template_bytes=None: rendered.update(data) or b"PK-docx"
        try:
            response = routes_contracts.get_contract_document("2004/BK-2026", db, None)
        finally:
            routes_contracts.doc_generator.render_contract_document = original_renderer

        self.assertEqual(response.body, b"PK-docx")
        self.assertEqual(rendered["customer_name"], "Tên hiện tại")
        self.assertEqual(rendered["customer_email"], "")
        self.assertEqual(rendered["due_date"], "")

    def test_document_route_uses_the_contracts_selected_private_template(self):
        current_rows = {
            "Contract": SimpleNamespace(
                id="2004/BK-2026", customer_id="customer-1", lead_id=None,
                contract_template_id="do-dac-v1", service_type="Đo hiện trạng",
                total_value=18500000, date_signed="2026-08-14",
            ),
            "Customer": SimpleNamespace(full_name="Tên hiện tại", phone="", address="", email=""),
            "ServiceLine": None,
            "Receivable": None,
            "LeadPipeline": None,
            "ContractGeneratedDocument": None,
            "ContractTemplate": [
                SimpleNamespace(
                    id="other-v2",
                    status="published",
                    version=2,
                    template_storage_key="contract-templates/other/v2.docx",
                ),
                SimpleNamespace(
                    id="do-dac-v1",
                    status="published",
                    version=1,
                    template_storage_key="contract-templates/MAU_HOP_DONG_DO_DAC_BACH_KHOA/v1.docx",
                ),
            ],
        }

        class CurrentDataQuery:
            def __init__(self, rows):
                self.rows = rows if isinstance(rows, list) else [rows]

            def filter(self, *criteria):
                for criterion in criteria:
                    key = getattr(getattr(criterion, "left", None), "key", None)
                    value = getattr(getattr(criterion, "right", None), "value", None)
                    if key is not None:
                        self.rows = [row for row in self.rows if getattr(row, key, None) == value]
                return self

            def order_by(self, *_ordering):
                return self

            def first(self):
                return self.rows[0] if self.rows else None

        db = SimpleNamespace(query=lambda model: CurrentDataQuery(current_rows.get(model.__name__)))
        rendered = {}
        requested_template_keys = []
        original_renderer = routes_contracts.doc_generator.render_contract_document
        original_reader = getattr(routes_contracts, "get_contract_template", None)
        routes_contracts.doc_generator.render_contract_document = lambda data, version, template_bytes=None: rendered.update(template_bytes=template_bytes) or b"PK-docx"
        routes_contracts.get_contract_template = lambda key: requested_template_keys.append(key) or b"PK-private-template"
        try:
            response = routes_contracts.get_contract_document("2004/BK-2026", db, None)
        finally:
            routes_contracts.doc_generator.render_contract_document = original_renderer
            if original_reader is None:
                del routes_contracts.get_contract_template
            else:
                routes_contracts.get_contract_template = original_reader

        self.assertEqual(response.body, b"PK-docx")
        self.assertEqual(rendered["template_bytes"], b"PK-private-template")
        self.assertEqual(
            requested_template_keys,
            ["contract-templates/MAU_HOP_DONG_DO_DAC_BACH_KHOA/v1.docx"],
        )

    def test_document_route_uses_selected_archived_template(self):
        current_rows = {
            "Contract": SimpleNamespace(
                id="2004/BK-2026", customer_id="customer-1", lead_id=None,
                contract_template_id="archived-v1", service_type="Đo hiện trạng",
                total_value=18500000, date_signed="2026-08-14",
            ),
            "Customer": SimpleNamespace(full_name="Tên hiện tại", phone="", address="", email=""),
            "ServiceLine": None,
            "Receivable": None,
            "LeadPipeline": None,
            "ContractGeneratedDocument": None,
            "ContractTemplate": [
                SimpleNamespace(
                    id="published-v2",
                    status="published",
                    version=2,
                    template_storage_key="contract-templates/current/v2.docx",
                ),
                SimpleNamespace(
                    id="archived-v1",
                    status="archived",
                    version=1,
                    template_storage_key="contract-templates/archived/v1.docx",
                ),
            ],
        }

        class CurrentDataQuery:
            def __init__(self, rows):
                self.rows = rows if isinstance(rows, list) else [rows]

            def filter(self, *criteria):
                for criterion in criteria:
                    key = getattr(getattr(criterion, "left", None), "key", None)
                    value = getattr(getattr(criterion, "right", None), "value", None)
                    if key is not None:
                        self.rows = [row for row in self.rows if getattr(row, key, None) == value]
                return self

            def order_by(self, *_ordering):
                return self

            def first(self):
                return self.rows[0] if self.rows else None

        db = SimpleNamespace(query=lambda model: CurrentDataQuery(current_rows.get(model.__name__)))
        requested_template_keys = []
        original_renderer = routes_contracts.doc_generator.render_contract_document
        original_reader = routes_contracts.get_contract_template
        routes_contracts.doc_generator.render_contract_document = lambda _data, _version, template_bytes=None: b"PK-docx"
        routes_contracts.get_contract_template = lambda key: requested_template_keys.append(key) or b"PK-private-template"
        try:
            response = routes_contracts.get_contract_document("2004/BK-2026", db, None)
        finally:
            routes_contracts.doc_generator.render_contract_document = original_renderer
            routes_contracts.get_contract_template = original_reader

        self.assertEqual(response.body, b"PK-docx")
        self.assertEqual(requested_template_keys, ["contract-templates/archived/v1.docx"])

    def test_renders_docx_bytes_without_creating_generated_docs_directory(self):
        """A new contract document must remain in memory, not in static/generated_docs."""
        with tempfile.TemporaryDirectory() as temp_dir:
            backend_dir = Path(temp_dir)
            template_dir = backend_dir / "src" / "templates"
            template_dir.mkdir(parents=True)
            template = Document()
            template.add_paragraph("Mã hợp đồng: {{contract_id}}")
            template.save(template_dir / "mau_hop_dong.docx")

            original_backend_dir = doc_generator.BACKEND_DIR
            doc_generator.BACKEND_DIR = str(backend_dir)
            try:
                output = doc_generator.render_contract_document(
                    {"contract_id": "2004/BK-2026"},
                    "mau_hop_dong_v1",
                )
            finally:
                doc_generator.BACKEND_DIR = original_backend_dir

            self.assertTrue(output.startswith(b"PK"))
            self.assertIn("2004/BK-2026", "".join(p.text for p in Document(io.BytesIO(output)).paragraphs))
            self.assertFalse((backend_dir / "static" / "generated_docs").exists())

    def test_renders_private_template_bytes_without_reading_repository_template(self):
        template = Document()
        template.add_paragraph("Khách hàng: {{customer_name}}")
        template_bytes = io.BytesIO()
        template.save(template_bytes)

        output = doc_generator.render_contract_document(
            {"customer_name": "Nguyễn Thị A"},
            "mau_hop_dong_v1",
            template_bytes=template_bytes.getvalue(),
        )

        self.assertIn("Nguyễn Thị A", "".join(p.text for p in Document(io.BytesIO(output)).paragraphs))


if __name__ == "__main__":
    unittest.main()
