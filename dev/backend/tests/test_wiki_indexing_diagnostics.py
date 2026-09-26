"""AI "học" tài liệu Wiki: lỗi phải có lý do hiển thị được, không im lặng."""
import io
from unittest.mock import MagicMock

import pytest
from docx import Document

from src.services import wiki_rag_service as rag


def _docx_bytes(paragraph: str, cells: list[str]) -> bytes:
    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=len(cells))
    for cell, value in zip(table.rows[0].cells, cells):
        cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_extraction_includes_table_cells():
    text = rag.extract_text_from_file(
        _docx_bytes("Điều 1. Phạm vi", ["Hành vi cấm", "Nhận quà của khách hàng"]),
        "ISO-001.docx",
    )
    assert "Điều 1. Phạm vi" in text
    assert "Hành vi cấm | Nhận quà của khách hàng" in text


@pytest.mark.parametrize(
    ("filename", "content", "reason"),
    [
        ("quy-che.doc", b"old word", "chưa hỗ trợ"),
        ("scan.txt", b"   \n  ", "Không đọc được chữ"),
    ],
)
def test_index_document_explains_why_nothing_was_indexed(filename, content, reason):
    with pytest.raises(ValueError, match=reason):
        rag.index_document(content, filename, "ISO-001", MagicMock())


def test_index_document_requires_a_gemini_key(monkeypatch):
    monkeypatch.setattr(rag, "_get_gemini_api_key", lambda: "")
    with pytest.raises(ValueError, match="Gemini API Key"):
        rag.index_document(b"Quy tac dao duc", "iso.txt", "ISO-001", MagicMock())


def test_index_document_returns_saved_chunk_count(monkeypatch):
    monkeypatch.setattr(rag, "_get_gemini_api_key", lambda: "AIza-test")
    monkeypatch.setattr(rag, "_embed_text", lambda *args, **kwargs: [0.1] * 3072)
    db = MagicMock()

    saved = rag.index_document(b"Quy tac dao duc nghe nghiep", "iso.txt", "ISO-001", db)

    assert saved == 1
    assert db.add.call_count == 1


def test_settings_key_wins_over_environment(monkeypatch):
    row = MagicMock(value=" AIza-from-settings ")
    session = MagicMock()
    session.query.return_value.filter.return_value.first.return_value = row
    monkeypatch.setattr(rag, "SessionLocal", lambda: session)
    monkeypatch.setenv("GEMINI_API_KEY", "AIza-stale-env")
    assert rag._get_gemini_api_key() == "AIza-from-settings"

    session.query.return_value.filter.return_value.first.return_value = None
    assert rag._get_gemini_api_key() == "AIza-stale-env"

    monkeypatch.setenv("GEMINI_API_KEY", "'your_gemini_api_key'")
    assert rag._get_gemini_api_key() == ""


def test_indexing_status_reports_chunks_progress_and_errors(monkeypatch):
    counts = [("DOC-OK", 12)]
    db = MagicMock()
    db.query.return_value.filter.return_value.group_by.return_value.all.return_value = counts
    started = []
    monkeypatch.setattr(rag, "_ensure_worker_running", lambda: started.append(True))
    monkeypatch.setattr(rag, "INDEXING_JOBS", {
        "DOC-RUN": {"status": "PROCESSING"},
        "DOC-BAD": {"status": "FAILED", "error": "Không đọc được chữ trong tài liệu"},
    })

    status = rag.get_indexing_status(["DOC-OK", "DOC-RUN", "DOC-BAD", "DOC-NEW"], db)

    assert status["DOC-OK"] == {"chunks": 12, "status": "COMPLETED", "error": None}
    assert status["DOC-RUN"]["status"] == "PROCESSING"
    assert status["DOC-BAD"] == {"chunks": 0, "status": "FAILED", "error": "Không đọc được chữ trong tài liệu"}
    assert status["DOC-NEW"]["status"] == "PENDING"
    # Worker không chạy sau khi khởi động lại thì tài liệu chờ mãi: xem trạng thái phải bật nó.
    assert started == [True]


def test_reset_indexing_job_clears_retry_history(monkeypatch):
    monkeypatch.setattr(rag, "INDEXING_JOBS", {"DOC": {"status": "FAILED", "retry_count": 3}})
    rag.reset_indexing_job("DOC")
    assert "DOC" not in rag.INDEXING_JOBS
